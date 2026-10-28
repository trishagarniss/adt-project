"""Membangun laporan DOCX Proyek 1 ADT dari output analisis.

Jalankan setelah `python src/analysis.py`:
    python src/build_report.py

Output:
    outputs/Laporan_PM25_Surakarta.docx
"""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
TAB = OUT / "tables"
FIG = OUT / "figures"
DOCX_PATH = OUT / "Laporan_PM25_Surakarta.docx"

TNR = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
TABLE_HEADER_FILL = "D9E2F3"
ACCESS_DATE = "29 Oktober 2026"

TABLES = 0
FIGURES = 0


def load_json() -> dict:
    return json.loads((OUT / "results.json").read_text(encoding="utf-8"))


def read_table(name: str) -> pd.DataFrame:
    return pd.read_csv(TAB / name)


def num(value, digits: int = 2) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "—"
    try:
        text = f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return str(value)
    return text.replace(".", ",")


def pval(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "—"
    value = float(value)
    if value < 0.001:
        return "<0,001"
    return f"{value:.3f}".replace(".", ",")


def tanggal(value: str) -> str:
    return pd.to_datetime(value).strftime("%d-%m-%Y")


def stat_value(value) -> str:
    text = str(value).strip()
    try:
        number = float(text)
    except ValueError:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return tanggal(text)
        return text
    if number.is_integer():
        return str(int(number))
    return num(number)


def param_value(value) -> str:
    text = str(value).strip()
    try:
        number = float(text)
    except ValueError:
        return text
    if number == 0:
        return "0"
    if abs(number) < 1e-4:
        return f"{number:.2e}".replace(".", ",")
    return num(number, 4)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(table, top=40, bottom=40, left=70, right=70) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def add_page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.name = TNR
    run.font.size = Pt(10)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def setup_document() -> Document:
    document = Document()
    document.core_properties.title = "Laporan Proyek 1 ADT: Peramalan PM2.5 Surakarta"
    document.core_properties.subject = "Analisis Data Time Series"
    document.core_properties.author = "Trisha Garnis"
    document.core_properties.keywords = "PM2.5, SARIMA, time series, Surakarta"
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    add_page_number_footer(section)

    normal = document.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), TNR)

    for name, size, before, after in (
        ("Heading 1", 14, 18, 8),
        ("Heading 2", 13, 12, 4),
        ("Heading 3", 12, 10, 3),
    ):
        style = document.styles[name]
        style.font.name = TNR
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True

    code_style = document.styles.add_style("Kode", 1)
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(7.5)
    code_style.paragraph_format.line_spacing = 1.0
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)

    return document


def add_rich(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def body(document: Document, text: str, first_line: bool = True, align=None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.first_line_indent = Cm(1.25) if first_line else Cm(0)
    if align is not None:
        paragraph.alignment = align
    add_rich(paragraph, text)


def bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(2)
    add_rich(paragraph, text)


def heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)


def caption(document: Document, text: str, kind: str = "Tabel") -> None:
    global TABLES, FIGURES
    if kind == "Tabel":
        TABLES += 1
        number = TABLES
    else:
        FIGURES += 1
        number = FIGURES
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"{kind} {number}. {text}")
    run.font.name = TNR
    run.font.size = Pt(10)
    run.bold = True


def source(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(f"Sumber: {text}")
    run.font.name = TNR
    run.font.size = Pt(9)
    run.italic = True


def add_table(
    document: Document,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    source_text: str,
    widths: list[float] | None = None,
    font_size: float = 9,
) -> None:
    caption(document, title, "Tabel")
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_margins(table)

    header_cells = table.rows[0].cells
    for index, text in enumerate(headers):
        cell = header_cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(text))
        run.font.name = TNR
        run.font.size = Pt(font_size)
        run.bold = True
        set_cell_shading(cell, TABLE_HEADER_FILL)
        if widths:
            cell.width = Cm(widths[index])
    set_repeat_table_header(table.rows[0])

    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            run.font.name = TNR
            run.font.size = Pt(font_size)
            if widths:
                cell.width = Cm(widths[index])

    source(document, source_text)


def add_figure(
    document: Document,
    path: Path,
    title: str,
    source_text: str,
    width_cm: float = 15.5,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    caption(document, title, "Gambar")
    source(document, source_text)


def add_code_listing(document: Document, path: Path, title: str) -> None:
    heading(document, title, 2)
    lines = path.read_text(encoding="utf-8").splitlines()
    for line in lines:
        paragraph = document.add_paragraph(style="Kode")
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(7.5)
    document.add_page_break()


def cover(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("LAPORAN PROYEK 1")
    run.font.name = TNR
    run.font.size = Pt(16)
    run.bold = True

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PENERAPAN MODEL MA, EXPONENTIAL SMOOTHING, AR, ARMA, ARIMA, DAN SARIMA")
    run.font.name = TNR
    run.font.size = Pt(14)
    run.bold = True

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PADA DATA DERET WAKTU PM2,5 KOTA SURAKARTA")
    run.font.name = TNR
    run.font.size = Pt(14)
    run.bold = True

    document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Mata Kuliah: Analisis Data Time Series (3 SKS)")
    run.font.name = TNR
    run.font.size = Pt(12)

    for _ in range(3):
        document.add_paragraph()

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Disusun oleh:")
    run.font.name = TNR
    run.font.size = Pt(12)

    for index in range(1, 6):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"[Nama Anggota {index}] – [NIM Anggota {index}]")
        run.font.name = TNR
        run.font.size = Pt(12)

    for _ in range(3):
        document.add_paragraph()

    for text in (
        "PROGRAM STUDI SAINS DATA",
        "SEMESTER 5",
        "2026",
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.name = TNR
        run.font.size = Pt(12)
        run.bold = True

    document.add_page_break()


def ringkasan(document: Document, results: dict) -> None:
    heading(document, "RINGKASAN", 1)
    text = (
        "Penelitian ini menganalisis konsentrasi PM2,5 mingguan di Kota Surakarta menggunakan data "
        "model CAMS global melalui Open-Meteo Air Quality API pada 8 Agustus 2022–21 Oktober 2024. "
        "Sebanyak 116 minggu homogen tanpa nilai hilang dipakai; 104 minggu menjadi data latih dan "
        "12 minggu terakhir data uji. Periode setelah Oktober 2024 dikeluarkan karena patahan rasio "
        "PM10/PM2,5. Penelitian menerapkan rata-rata bergerak, SES, Holt, Holt-Winters, AR, ARMA, "
        "ARIMA, SARIMA, dan Fourier-ARIMA. Stasioneritas diuji dengan ADF dan KPSS; pemilihan model "
        "memakai AICc dan Ljung-Box; evaluasi memakai MAE, RMSE, MAPE, dan MASE. Model terpilih "
        f"{results['final_model']} dengan RMSE uji {num(results['test_rmse'])} µg/m³, MAE "
        f"{num(results['test_mae'])} µg/m³, MAPE {num(results['test_mape'])}%, dan MASE "
        f"{num(results['test_mase'], 3)}. Namun naive random walk menghasilkan RMSE 4,64 dan uji "
        "Diebold-Mariano tidak signifikan (p=0,197), sehingga keunggulan model tidak dapat "
        "dibedakan. Peramalan 12 minggu relatif datar sekitar 14–15 µg/m³ dengan interval 80% dan "
        "95% yang melebar. Keterbatasan utama adalah panjang seri hanya 2,2 siklus tahunan, suku "
        "musiman tidak signifikan, dan residual heteroskedastik."
    )
    body(document, text, first_line=False)
    document.add_page_break()


def bab1(document: Document) -> None:
    heading(document, "BAB 1 PENDAHULUAN", 1)

    heading(document, "1.1 Latar Belakang", 2)
    body(
        document,
        "PM2,5 merupakan partikel halus berdiameter aerodinamis tidak lebih dari 2,5 mikrometer "
        "yang dapat masuk jauh ke saluran pernapasan dan berkontribusi terhadap penyakit "
        "kardiovaskular, pernapasan, serta kematian dini (Cohen et al., 2017). Sebagai kota dengan "
        "aktivitas transportasi, perdagangan, dan permukiman yang padat, Surakarta menghadapi "
        "tekanan kualitas udara yang perlu dipantau secara berkala. Peramalan konsentrasi PM2,5 "
        "bermanfaat untuk memberi peringatan dini, merancang kebijakan, dan mengalokasikan sumber "
        "daya kesehatan masyarakat.",
    )
    body(
        document,
        "Ketersediaan stasiun pengukuran PM2,5 darat di Surakarta terbatas. Oleh karena itu "
        "digunakan data sekunder dari model CAMS global yang disediakan melalui Open-Meteo Air "
        "Quality API. Data ini merupakan keluaran model reanalisis, bukan pengukuran stasiun, "
        "sehingga hasil harus ditafsirkan sebagai representasi grid atmosfer di sekitar koordinat "
        "Surakarta (Open-Meteo, 2026; Copernicus Atmosphere Monitoring Service, 2024).",
    )
    body(
        document,
        "Pemeriksaan awal menemukan patahan struktural pada November 2024: rasio PM10/PM2,5 turun "
        "dari sekitar 1,46 menjadi 1,04. Karena rasio tersebut tidak lagi masuk akal secara fisik, "
        "periode setelah Oktober 2024 dikeluarkan dari pemodelan. Keputusan ini menjaga homogenitas "
        "seri dan menjadi bagian dari kontrol kualitas data.",
    )

    heading(document, "1.2 Tujuan Analisis", 2)
    for text in (
        "Menyiapkan dan mendeskripsikan seri mingguan PM2,5 Kota Surakarta dari sumber sekunder yang dapat ditelusuri.",
        "Mengeksplorasi pola tren, musiman, autokorelasi, dan stasioneritas seri.",
        "Mengestimasi dan membandingkan model MA, Exponential Smoothing, AR, ARMA, ARIMA, dan SARIMA.",
        "Mendiagnosis residual serta memilih model dengan kriteria yang valid.",
        "Meramalkan PM2,5 untuk 12 minggu ke depan beserta interval ketidakpastian 80% dan 95%.",
    ):
        bullet(document, text)

    heading(document, "1.3 Manfaat", 2)
    body(
        document,
        "Secara praktis, hasil peramalan dapat dipakai sebagai skenario awal untuk peringatan dini "
        "dan komunikasi risiko. Secara akademik, laporan ini mendokumentasikan alur analisis deret "
        "waktu secara lengkap, termasuk keterbatasan data dan asumsi yang dilanggar, sehingga dapat "
        "dijalankan ulang dan dikritisi.",
    )


def bab2(document: Document, results: dict, ratios: pd.DataFrame) -> None:
    heading(document, "BAB 2 DATA DAN EKSPLORASI", 1)

    heading(document, "2.1 Sumber, Satuan, dan Periode", 2)
    body(
        document,
        "Data diunduh dari Open-Meteo Air Quality API yang meneruskan keluaran model CAMS global. "
        "Koordinat yang dipakai adalah 7,5545° LS dan 110,8060° BT dengan zona waktu Asia/Jakarta. "
        "Variabel yang diambil adalah pm2_5 dalam satuan µg/m³. Data mentah berbentuk jam-jaman, "
        "kemudian diagregasi menjadi rata-rata mingguan. Agregasi dilakukan oleh penulis, bukan "
        "oleh penyedia, sehingga bentuk data mentah tetap dapat ditelusuri. Tanggal akses data: "
        f"{ACCESS_DATE}. Ringkasan data dan pembagian latih–uji disajikan pada Tabel 1.",
    )
    body(
        document,
        "Tautan sumber: https://open-meteo.com/en/docs/air-quality-api dan "
        "https://ads.atmosphere.copernicus.eu/datasets/cams-global-reanalysis-eac4. Data CAMS "
        "digunakan sebagai representasi kondisi atmosfer; validasi terhadap stasiun darat "
        "Surakarta tidak tersedia dalam proyek ini.",
    )

    raw = pd.read_csv(ROOT / "data" / "raw" / "openmeteo_pm25_surakarta_hourly.csv")
    raw["pm2_5"] = pd.to_numeric(raw["pm2_5"], errors="coerce")
    valid = raw["pm2_5"].notna().sum()
    missing = int(raw["pm2_5"].isna().sum())
    add_table(
        document,
        "Ringkasan data dan pembagian latih–uji",
        ["Keterangan", "Nilai"],
        [
            ["Periode model", f"{tanggal(results['start'])} – {tanggal(results['end'])}"],
            ["Frekuensi", "Mingguan (rata-rata 168 jam)"],
            ["Jumlah observasi", str(results["n"])],
            ["Data latih", f"{results['n_train']} minggu"],
            ["Data uji", f"{results['n_test']} minggu"],
            ["Jam mentah", f"{len(raw):,}".replace(",", ".")],
            ["Jam valid", f"{valid:,}".replace(",", ".")],
            ["Jam hilang", f"{missing} (di awal seri)"],
            ["Nilai hilang pada seri model", "0"],
        ],
        "Open-Meteo/CAMS (2026), diolah penulis.",
        widths=[7.0, 7.0],
    )
    body(
        document,
        "Dari 37.200 jam mentah, 37.121 jam valid dan 79 jam hilang. Seluruh nilai hilang berada di "
        "awal periode (1–4 Agustus 2022) dan tidak berada di dalam rentang model. Setelah agregasi "
        "dan penyaringan kualitas, terbentuk 116 minggu lengkap tanpa nilai hilang.",
    )

    heading(document, "2.2 Patahan Struktural November 2024", 2)
    body(
        document,
        "Rasio PM10/PM2,5 bulanan dipakai sebagai indikator konsistensi fisik. Sebelum November "
        "2024, rasio rata-rata 1,464 dengan simpangan baku 0,015; sesudahnya rasio turun menjadi "
        "1,036 dengan simpangan baku 0,025. Nilai mendekati 1 berarti PM2,5 hampir sama dengan "
        "PM10, yang tidak wajar untuk aerosol atmosfer. Pola serupa ditemukan pada lokasi lain "
        "(Jakarta dan Surabaya), sehingga indikasi patahan model CAMS lebih kuat daripada sekadar "
        "fluktuasi lokal. Periode setelah 31 Oktober 2024 karena itu tidak dipakai (Gambar 1; Tabel 2).",
    )
    add_figure(
        document,
        FIG / "fig01_data_break.png",
        "Rasio PM10/PM2,5 bulanan dan batas periode homogen",
        "Open-Meteo/CAMS (2026), diolah penulis.",
    )
    pre = ratios.loc[ratios["homogeneous"], "ratio"]
    post = ratios.loc[~ratios["homogeneous"], "ratio"]
    add_table(
        document,
        "Ringkasan rasio PM10/PM2,5 sebelum dan sesudah patahan",
        ["Periode", "n bulan", "Rata-rata", "Simpangan baku", "Status"],
        [
            ["Sebelum Nov 2024", str(len(pre)), num(pre.mean(), 3), num(pre.std(), 3), "Homogen"],
            ["Nov 2024 dan sesudahnya", str(len(post)), num(post.mean(), 3), num(post.std(), 3), "Dikeluarkan"],
        ],
        "Open-Meteo/CAMS (2026), diolah penulis.",
        widths=[4.2, 1.8, 2.3, 2.8, 2.5],
    )

    heading(document, "2.3 Statistik Deskriptif", 2)
    desc = read_table("descriptive.csv")
    body(
        document,
        "Rata-rata PM2,5 mingguan adalah 23,82 µg/m³ dengan simpangan baku 8,45 µg/m³. Nilai "
        "minimum 9,69 µg/m³ dan maksimum 47,26 µg/m³. Distribusi sedikit miring ke kanan "
        "(kemencengan 0,635) dengan ekor relatif ringan (kurtosis −0,367). Koefisien variasi 0,355 "
        "menunjukkan variasi yang cukup besar terhadap rata-rata (Tabel 3). Pola seri dan "
        "pembagian latih–uji ditunjukkan pada Gambar 2.",
    )
    add_table(
        document,
        "Statistik deskriptif seri mingguan PM2,5",
        ["Statistik", "Nilai"],
        [[row["Statistik"], stat_value(row["Nilai"])] for _, row in desc.iterrows()],
        "Open-Meteo/CAMS (2026), diolah penulis.",
        widths=[7.0, 7.0],
        font_size=9,
    )
    add_figure(
        document,
        FIG / "fig02_series.png",
        "Seri mingguan PM2,5 Kota Surakarta dan pembagian data latih–uji",
        "Open-Meteo/CAMS (2026), diolah penulis.",
    )

    heading(document, "2.4 Dekomposisi dan Pola Musiman", 2)
    body(
        document,
        "Dekomposisi STL dilakukan dengan periode 52 minggu, seasonal smoother 7, dan sifat robust. "
        "Model aditif dipilih karena amplitudo musiman relatif stabil terhadap level dan seri selalu "
        "positif. Kekuatan musiman yang dihitung dari rasio varians komponen musiman terhadap "
        f"seri terdetrend adalah {num(results['seasonal_strength'], 3)}. Nilai ini tinggi, tetapi "
        "harus dibaca hati-hati karena panjang seri hanya sekitar 2,2 siklus tahunan (Gambar 3).",
    )
    add_figure(
        document,
        FIG / "fig03_decomposition.png",
        "Dekomposisi STL seri mingguan (periode 52, seasonal smoother 7)",
        "Open-Meteo/CAMS (2026), diolah penulis.",
    )
    body(
        document,
        "Profil bulanan menunjukkan rata-rata lebih tinggi pada Januari–Maret dan lebih rendah pada "
        "Juli–September. Pola ini bersifat deskriptif karena setiap bulan hanya terwakili dua sampai "
        "tiga kali (Gambar 4). Estimasi komponen musiman tahunan dengan demikian rapuh dan tidak "
        "boleh dijadikan satu-satunya dasar kebijakan.",
    )
    add_figure(
        document,
        FIG / "fig04_seasonal_profile.png",
        "Profil musiman rata-rata menurut bulan kalender",
        "Open-Meteo/CAMS (2026), diolah penulis.",
    )

    heading(document, "2.5 ACF dan PACF", 2)
    body(
        document,
        "ACF level meluruh lambat, mengindikasikan ketidakstasioneran. Setelah diferensiasi pertama, "
        "ACF menghilang lebih cepat. Pada lag musiman 52 masih terlihat pola, tetapi jumlah siklus "
        "yang tersedia terlalu sedikit untuk mengestimasi struktur musiman panjang secara stabil. "
        "Karena itu, model musiman diuji sebagai kandidat, bukan diasumsikan pasti benar (Gambar 5).",
    )
    add_figure(
        document,
        FIG / "fig05_acf_pacf.png",
        "ACF dan PACF level serta diferensiasi pertama (lag 1–55)",
        "Open-Meteo/CAMS (2026), diolah penulis.",
    )


def bab3(document: Document) -> None:
    heading(document, "BAB 3 METODE", 1)

    heading(document, "3.1 Rata-rata Bergerak dan Exponential Smoothing", 2)
    body(
        document,
        "Rata-rata bergerak sederhana dengan jendela 3 dan 8 minggu dipakai sebagai penghalus dan "
        "pembanding naif. Jendela 3 minggu dipilih untuk menangkap perubahan jangka pendek, "
        "sedangkan jendela 8 minggu memberi penghalusan lebih kuat. Keduanya bukan model "
        "probabilistik sehingga tidak memiliki AIC yang sebanding dengan keluarga ARIMA.",
    )
    body(
        document,
        "SES meramalkan level dengan pembaruan ℓₜ = αyₜ + (1−α)ℓₜ₋₁. Holt menambahkan tren "
        "bₜ = β(ℓₜ−ℓₜ₋₁) + (1−β)bₜ₋₁. Holt-Winters menambahkan komponen musiman aditif dengan "
        "periode 52. Parameter smoothing dan nilai awal diestimasi dengan maximum likelihood. "
        "Komponen tren dan musiman hanya dipertahankan bila memberi perbaikan yang dapat "
        "dijelaskan; AIC antar-ES dan ARIMA tidak dibandingkan langsung (Hyndman & Athanasopoulos, "
        "2021).",
    )

    heading(document, "3.2 AR, MA, ARMA, dan ARIMA", 2)
    body(
        document,
        "Model AR(p) menyatakan yₜ sebagai fungsi linear dari p nilai sebelumnya; MA(q) menyatakan "
        "yₜ sebagai fungsi linear dari q galat sebelumnya; ARMA(p,q) menggabungkan keduanya. "
        "Stasioneritas AR diperiksa dari akar polinomial AR, sedangkan invertibilitas MA diperiksa "
        "dari akar polinomial MA. Semua akar harus berada di luar lingkaran satuan (Box et al., "
        "2015; Shumway & Stoffer, 2017).",
    )
    body(
        document,
        "ARIMA(p,d,q) menambahkan diferensiasi ke-d. Nilai d ditentukan dari ADF/KPSS dan ACF, bukan "
        "dari pengulangan differencing sampai grafik tampak rata. Kandidat ARIMA dicari dengan grid "
        "sistematis p,q = 0–3 pada d = 1. Pencarian ini berfungsi sebagai pencarian otomatis; "
        "hasilnya dikritik dengan membandingkan AICc, Ljung-Box, dan RMSE uji. `pmdarima` tidak "
        "dipakai karena lingkungan memakai NumPy 2.4.6 yang tidak kompatibel dengan versi "
        "`pmdarima` yang tersedia.",
    )

    heading(document, "3.3 SARIMA dan Fourier", 2)
    body(
        document,
        "SARIMA(p,d,q)(P,D,Q)ₘ menambahkan operator musiman. Dengan m = 52 dan hanya 116 minggu, "
        "D = 1 akan menghabiskan terlalu banyak observasi; karena itu D = 0 dan struktur musiman "
        "ditangkap melalui P,Q. Kandidat dibatasi p,q,P,Q ∈ {0,1,2} dan disaring dengan "
        "konvergensi serta Ljung-Box lag 20. Model dengan AICc terendah yang gagal Ljung-Box tidak "
        "dipilih.",
    )
    body(
        document,
        "Sebagai bonus, regresi Fourier dengan K = 1,2,3 harmonik ditambahkan pada ARIMA(1,1,0), "
        "ARIMA(1,1,1), dan ARIMA(2,1,0). Regresor Fourier fleksibel, tetapi dapat menghasilkan "
        "koefisien MA yang mendekati batas non-invertibilitas; hasilnya dilaporkan sebagai "
        "pembanding, bukan model utama.",
    )

    heading(document, "3.4 Kriteria Evaluasi", 2)
    body(
        document,
        "Data uji adalah 12 minggu terakhir dan tidak dipakai dalam identifikasi maupun estimasi. "
        "Akurasi diukur dengan MAE, RMSE, MAPE, sMAPE, dan MASE. MASE memakai seasonal naive lag 52 "
        "pada data latih sebagai skala. AIC/AICc/BIC hanya dibandingkan dalam keluarga ARIMA dengan "
        "diferensiasi dan data yang sama. Perbedaan akurasi antar-model diuji dengan Diebold-Mariano. "
        "Residual diperiksa dengan ACF, Ljung-Box, Jarque-Bera, Shapiro-Wilk, ARCH LM, dan "
        "Durbin-Watson. Interval peramalan disajikan pada 80% dan 95%. Diagram alur penelitian "
        "disajikan pada Gambar 6.",
    )
    add_figure(
        document,
        FIG / "fig00_flowchart.png",
        "Diagram alur penelitian",
        "Penulis (2026).",
        width_cm=15.0,
    )


def bab4(document: Document, results: dict) -> None:
    heading(document, "BAB 4 HASIL DAN PEMBAHASAN", 1)

    heading(document, "4.1 Stasioneritas dan Transformasi", 2)
    body(
        document,
        "Pada level, ADF tidak menolak akar unit (p = 0,203), sedangkan KPSS tidak menolak "
        "stasioneritas (p = 0,100). Konflik ini lazim pada sampel pendek. ACF level yang meluruh "
        "lambat dan hasil diferensiasi pertama yang sangat signifikan (ADF p < 0,001) mendukung "
        "d = 1. Diferensiasi musiman lag 52 juga stasioner, tetapi dengan hanya 63–64 observasi "
        "efektif dan 2,2 siklus, D = 1 berisiko over-differencing. Model final memakai d = 1, "
        "D = 0, dan komponen AR musiman (Tabel 4).",
    )
    body(
        document,
        "Transformasi log tidak dipakai karena varians level tidak menunjukkan pembesaran yang "
        "konsisten seiring naiknya level, dan transformasi akan menyulitkan interpretasi µg/m³. "
        "Heteroskedastisitas residual tetap dideteksi pada tahap diagnostik dan dibahas pada "
        "Bagian 4.6.",
    )
    stat = read_table("stationarity.csv")
    add_table(
        document,
        "Hasil uji ADF dan KPSS",
        ["Seri", "Uji", "Statistik", "p-value", "Lag", "n", "Kesimpulan"],
        [
            [
                row["Seri"],
                row["Uji"],
                num(row["Statistik"], 3),
                pval(row["p-value"]),
                str(int(row["Lag"])),
                str(int(row["n"])),
                row["Kesimpulan"],
            ]
            for _, row in stat.iterrows()
        ],
        "Hasil pengolahan penulis (2026).",
        widths=[3.4, 1.2, 1.8, 1.5, 1.0, 1.0, 3.0],
        font_size=8.5,
    )

    heading(document, "4.2 Moving Average dan Exponential Smoothing", 2)
    comp = read_table("model_comparison.csv")
    es_names = ["MA(3)", "MA(8)", "SES", "Holt", "Holt-Winters"]
    es_rows = []
    for _, row in comp[comp["Model"].isin(es_names)].iterrows():
        es_rows.append([
            row["Model"],
            row["Spesifikasi"],
            num(row["RMSE"], 3),
            num(row["MAE"], 3),
            num(row["MAPE"], 2),
            num(row["MASE"], 3),
        ])
    add_table(
        document,
        "Perbandingan moving average dan Exponential Smoothing pada data uji",
        ["Model", "Spesifikasi", "RMSE", "MAE", "MAPE (%)", "MASE"],
        es_rows,
        "Hasil pengolahan penulis (2026).",
        widths=[2.4, 5.0, 1.6, 1.6, 1.6, 1.5],
        font_size=8.5,
    )
    body(
        document,
        "MA(3) memberi RMSE 4,660, lebih baik daripada MA(8) (6,687). SES dan Holt menghasilkan "
        "RMSE sekitar 5,46 dan 5,36; keduanya lebih buruk daripada naive. Holt-Winters memiliki "
        "RMSE 6,739 dan gagal Ljung-Box (p = 0,006), sehingga tidak layak dipakai. Parameter "
        "Holt-Winters juga berada di batas (α ≈ 0, β = 0, γ = 0), tanda bahwa komponen musiman "
        "tidak teridentifikasi dengan baik pada seri sependek ini (Tabel 5; Tabel 6).",
    )
    es_params = read_table("es_params.csv")
    keep = es_params[es_params["Parameter"].isin(
        ["smoothing_level", "smoothing_trend", "smoothing_seasonal", "initial_level", "initial_trend", "initial_seasons"]
    )]
    rows = []
    for _, row in keep.iterrows():
        rows.append([row["Model"], row["Parameter"], param_value(row["Nilai"])])
    add_table(
        document,
        "Parameter smoothing dan nilai awal model Exponential Smoothing",
        ["Model", "Parameter", "Nilai"],
        rows,
        "Hasil pengolahan penulis (2026).",
        widths=[3.0, 4.0, 7.0],
        font_size=8.5,
    )

    heading(document, "4.3 AR, MA, dan ARMA", 2)
    ar = read_table("grid_ar.csv").head(5)
    arma = read_table("grid_arma.csv").head(5)
    rows = []
    for _, row in ar.iterrows():
        rows.append([f"AR({int(row['p'])})", "—", num(row["AIC"], 2), num(row["AICc"], 2), num(row["BIC"], 2)])
    for _, row in arma.iterrows():
        rows.append([
            f"ARMA({int(row['p'])},{int(row['q'])})",
            f"d={int(row['d'])}",
            num(row["AIC"], 2),
            num(row["AICc"], 2),
            num(row["BIC"], 2),
        ])
    add_table(
        document,
        "Kandidat AR dan ARMA teratas menurut AICc",
        ["Model", "Diferensiasi", "AIC", "AICc", "BIC"],
        rows,
        "Hasil pengolahan penulis (2026).",
        widths=[3.2, 2.2, 2.4, 2.4, 2.4],
        font_size=8.5,
    )
    body(
        document,
        "AR(12) terpilih dalam grid AR dengan AICc 646,33, tetapi RMSE ujinya 6,572. ARMA(2,3) "
        "terpilih dalam grid ARMA dengan AICc 688,05 dan RMSE 7,067 (Tabel 7). Keduanya stabil: "
        "modulus akar AR minimum 1,147 dan modulus akar MA minimum 1,124, sehingga syarat "
        "stasioneritas dan invertibilitas terpenuhi (Tabel 8). Meskipun demikian, akurasi keduanya "
        "jauh di bawah naive dan SARIMA, sehingga tidak dipilih sebagai model final.",
    )
    roots = read_table("roots.csv")
    add_table(
        document,
        "Pemeriksaan akar model AR, ARMA, dan SARIMA",
        ["Model", "Komponen", "Jumlah akar", "|root| minimum", "|root| maksimum", "Kesimpulan"],
        [
            [
                row["Model"],
                row["Komponen"],
                str(int(row["Jumlah akar"])),
                num(row["|root| minimum"], 4),
                num(row["|root| maksimum"], 4),
                row["Kesimpulan"],
            ]
            for _, row in roots.iterrows()
        ],
        "Hasil pengolahan penulis (2026).",
        widths=[4.0, 1.6, 1.3, 2.0, 2.0, 3.2],
        font_size=8,
    )

    heading(document, "4.4 ARIMA", 2)
    arima = read_table("grid_arima.csv").head(5)
    add_table(
        document,
        "Kandidat ARIMA teratas menurut AICc (d = 1)",
        ["p", "d", "q", "AIC", "AICc", "BIC", "Ljung-Box p", "Konvergen"],
        [
            [
                str(int(row["p"])),
                str(int(row["d"])),
                str(int(row["q"])),
                num(row["AIC"], 2),
                num(row["AICc"], 2),
                num(row["BIC"], 2),
                pval(row["LB20_p"]),
                "Ya" if bool(row["Konvergen"]) else "Tidak",
            ]
            for _, row in arima.iterrows()
        ],
        "Hasil pengolahan penulis (2026).",
        widths=[0.9, 0.9, 0.9, 2.1, 2.1, 2.1, 2.2, 1.8],
        font_size=8.5,
    )
    body(
        document,
        "ARIMA(1,1,3) terpilih dengan AICc 680,27 dan Ljung-Box p = 0,871, tetapi RMSE ujinya "
        "5,708. Kandidat ARIMA(0,1,3) dan ARIMA(2,1,3) memiliki AICc yang berdekatan, sehingga "
        "perbedaan di antara mereka tipis (Tabel 9). Pencarian otomatis ini menunjukkan bahwa kriteria "
        "informasi saja tidak cukup: model ARIMA non-musiman tetap kalah dari naive pada data uji.",
    )

    heading(document, "4.5 SARIMA", 2)
    sarima = read_table("grid_sarima.csv").head(7)
    add_table(
        document,
        "Kandidat SARIMA teratas menurut AICc",
        ["p", "d", "q", "P", "D", "Q", "s", "AICc", "Ljung-Box p", "Konvergen"],
        [
            [
                str(int(row["p"])),
                str(int(row["d"])),
                str(int(row["q"])),
                str(int(row["P"])),
                str(int(row["D"])),
                str(int(row["Q"])),
                str(int(row["s"])),
                num(row["AICc"], 2),
                pval(row["LB20_p"]),
                "Ya" if bool(row["Konvergen"]) else "Tidak",
            ]
            for _, row in sarima.iterrows()
        ],
        "Hasil pengolahan penulis (2026).",
        widths=[0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 1.9, 2.0, 1.6],
        font_size=8,
    )
    body(
        document,
        "Tiga kandidat dengan AICc terendah memakai Q = 1 dan memiliki Ljung-Box p < 0,01, sehingga "
        "residualnya masih mengandung autokorelasi. Model SARIMA(2,1,0)(1,0,0)[52] berada pada "
        "peringkat berikutnya dengan AICc 310,63 dan Ljung-Box p = 0,081. Model ini dipilih karena "
        "merupakan kandidat ber-AICc rendah yang lolos penyaringan residual, bukan karena AICc "
        "terendah secara mutlak (Tabel 10). Kandidat dengan D = 1 diuji terpisah dan tidak memberi perbaikan "
        "yang cukup untuk mengimbangi hilangnya observasi.",
    )
    body(
        document,
        "Persamaan model final dapat dituliskan sebagai (1 − φ₁B − φ₂B²)(1 − Φ₁B⁵²)(1 − B)Yₜ = εₜ, "
        "dengan φ₁ = −0,317, φ₂ = −0,191, dan Φ₁ = 0,0737. Koefisien AR(1) signifikan pada α = 5% "
        "(p = 0,015), sedangkan AR(2) (p = 0,238) dan AR musiman (p = 0,409) tidak signifikan. "
        "Akibatnya, struktur musiman yang diestimasi lemah dan tidak boleh ditafsirkan sebagai "
        "bukti kuat adanya siklus tahunan.",
    )

    heading(document, "4.6 Diagnostik Residual", 2)
    diag = read_table("final_diagnostics.csv")
    add_table(
        document,
        "Diagnostik residual model final",
        ["Uji", "Statistik", "p-value", "Interpretasi"],
        [
            [row["Uji"], num(row["Statistik"], 3), pval(row["p_value"]), row["Interpretasi"]]
            for _, row in diag.iterrows()
        ],
        "Hasil pengolahan penulis (2026).",
        widths=[4.2, 2.4, 2.0, 5.2],
        font_size=8.5,
    )
    body(
        document,
        "Ljung-Box lag 10, 20, dan 52 tidak menolak tidak adanya autokorelasi (p = 0,308; 0,097; "
        "0,246). Jarque-Bera (p = 0,119) dan Shapiro-Wilk (p = 0,097) tidak menolak normalitas. "
        "Durbin-Watson 2,146 mendekati 2 (Tabel 11; Gambar 7). Namun ARCH LM menolak homoskedastisitas (p = 0,024), "
        "sehingga varians residual berubah-ubah. Konsekuensinya, interval peramalan yang dihitung "
        "dengan asumsi varians konstan cenderung terlalu sempit pada periode bergejolak dan terlalu "
        "lebar pada periode tenang.",
    )
    add_figure(
        document,
        FIG / "fig08_residual_diagnostics.png",
        "Diagnostik residual model SARIMA final",
        "Hasil pengolahan penulis (2026).",
    )

    heading(document, "4.7 Perbandingan Model", 2)
    rows = []
    for _, row in comp.iterrows():
        rows.append([
            row["Model"],
            num(row["AICc"], 2),
            num(row["RMSE"], 3),
            num(row["MAE"], 3),
            num(row["MAPE"], 2),
            num(row["MASE"], 3),
            pval(row["LB20_p"]) if row["AIC_sebanding"] == "Ya" else "—",
            row["AIC_sebanding"],
        ])
    add_table(
        document,
        "Perbandingan seluruh model pada data uji",
        ["Model", "AICc", "RMSE", "MAE", "MAPE (%)", "MASE", "LB20 p", "AIC sebanding"],
        rows,
        "Hasil pengolahan penulis (2026). AICc hanya sebanding untuk model ARIMA dengan d dan data yang sama.",
        widths=[4.2, 1.4, 1.4, 1.4, 1.4, 1.2, 1.2, 1.6],
        font_size=7.5,
    )
    body(
        document,
        "Naive random walk memiliki RMSE 4,643 dan MAE 2,769, sedikit lebih baik daripada SARIMA "
        "(RMSE 4,814; MAE 2,951). MA(3) berada di antara keduanya (RMSE 4,660). Perbedaan SARIMA "
        "dan naive tidak signifikan secara statistik (Diebold-Mariano p = 0,197), sehingga keduanya "
        "tidak dapat dibedakan secara meyakinkan. SARIMA tetap dipilih karena memberi struktur "
        "musiman dan interval peramalan eksplisit, bukan karena unggul akurasi (Tabel 12; "
        "Gambar 8; Gambar 9; Tabel 13).",
    )
    add_figure(
        document,
        FIG / "fig06_model_rmse.png",
        "Perbandingan RMSE model pada 12 minggu uji",
        "Hasil pengolahan penulis (2026).",
    )
    add_figure(
        document,
        FIG / "fig07_test_forecasts.png",
        "Aktual versus peramalan beberapa model pada periode uji",
        "Hasil pengolahan penulis (2026).",
    )
    dm = read_table("dm_tests.csv")
    dm_keep = dm[dm["Model pembanding"].isin(
        ["Naive (random walk)", "MA(3)", "Fourier K=2 + ARIMA(1,1,1)", "Holt", "ARIMA(1,1,3)"]
    )]
    add_table(
        document,
        "Uji Diebold-Mariano model final terhadap pembanding utama",
        ["Model pembanding", "DM statistik", "p-value", "Kesimpulan"],
        [
            [row["Model pembanding"], num(row["DM statistik"], 3), pval(row["p-value"]), row["Kesimpulan"]]
            for _, row in dm_keep.iterrows()
        ],
        "Hasil pengolahan penulis (2026).",
        widths=[5.0, 2.2, 2.0, 4.6],
        font_size=8.5,
    )

    heading(document, "4.8 Peramalan dan Interpretasi", 2)
    forecast = read_table("forecast.csv")
    add_table(
        document,
        "Peramalan PM2,5 12 minggu ke depan dengan interval 80% dan 95%",
        ["Minggu", "Peramalan", "Bawah 80%", "Atas 80%", "Bawah 95%", "Atas 95%"],
        [
            [
                tanggal(row["Minggu"]),
                num(row["Peramalan"], 2),
                num(row["Batas_bawah_80"], 2),
                num(row["Batas_atas_80"], 2),
                num(row["Batas_bawah_95"], 2),
                num(row["Batas_atas_95"], 2),
            ]
            for _, row in forecast.iterrows()
        ],
        "Hasil pengolahan penulis (2026); satuan µg/m³.",
        widths=[2.2, 2.2, 2.1, 2.1, 2.1, 2.1],
        font_size=8.5,
    )
    body(
        document,
        "Peramalan bergerak datar dari 14,77 µg/m³ pada minggu pertama menjadi 14,49 µg/m³ pada "
        "minggu terakhir, dengan rata-rata 14,55 µg/m³. Interval 80% melebar dari 8,02–21,52 "
        "menjadi −1,77–30,75; interval 95% melebar dari 4,45–25,09 menjadi −10,37–39,35. Batas "
        "bawah negatif adalah artefak aproksimasi normal dan tidak memiliki makna fisik. Secara "
        "substantif, model memprediksi stabilisasi pada tingkat sedang, tetapi ketidakpastiannya "
        "besar sehingga angka ini lebih tepat dibaca sebagai skenario (Tabel 14; Gambar 10).",
    )
    add_figure(
        document,
        FIG / "fig09_forecast.png",
        "Peramalan PM2,5 12 minggu ke depan dengan interval 80% dan 95%",
        "Hasil pengolahan penulis (2026).",
    )


def bab5(document: Document, results: dict) -> None:
    heading(document, "BAB 5 KESIMPULAN", 1)
    heading(document, "5.1 Kesimpulan", 2)
    for text in (
        f"Data PM2,5 mingguan Kota Surakarta yang homogen mencakup 116 minggu (8 Agustus 2022–21 Oktober 2024) tanpa nilai hilang; 104 minggu dipakai untuk latih dan 12 minggu untuk uji.",
        "Patahan rasio PM10/PM2,5 pada November 2024 terdeteksi dan periode sesudahnya dikeluarkan agar seri tetap konsisten secara fisik.",
        "Uji ADF/KPSS dan ACF/PACF mendukung diferensiasi pertama (d = 1); D = 0 dipilih untuk menghindari over-differencing musiman.",
        f"Model terpilih adalah {results['final_model']} dengan AICc latih 310,63 dan RMSE uji 4,814 µg/m³.",
        "Naive random walk memiliki RMSE 4,643; uji Diebold-Mariano tidak signifikan (p = 0,197), sehingga keunggulan SARIMA tidak dapat dibedakan secara meyakinkan.",
        "Peramalan 12 minggu relatif datar sekitar 14–15 µg/m³, dengan interval 80% dan 95% yang melebar cepat.",
    ):
        bullet(document, text)

    heading(document, "5.2 Keterbatasan", 2)
    for text in (
        "Data berasal dari grid model CAMS, bukan stasiun pengukuran darat Surakarta, sehingga tidak menangkap variasi lokal skala halus.",
        "Panjang seri hanya 116 minggu atau sekitar 2,2 siklus tahunan; estimasi komponen musiman rapuh dan suku AR musiman tidak signifikan.",
        "Residual model final heteroskedastik (ARCH LM p = 0,024), sehingga interval peramalan dengan varians konstan tidak sepenuhnya valid.",
        "Tidak ada validasi silang bergulir (rolling-origin); evaluasi hanya memakai satu periode uji 12 minggu.",
        "Model ARIMA/SARIMA mengasumsikan hubungan linear; potensi efek libur, kebakaran lahan, dan perubahan kebijakan belum dimasukkan sebagai variabel eksogen.",
    ):
        bullet(document, text)

    heading(document, "5.3 Saran", 2)
    body(
        document,
        "Untuk penelitian lanjutan, disarankan membandingkan data CAMS dengan stasiun darat bila "
        "tersedia, menambah variabel eksogen seperti curah hujan, kecepatan angin, dan titik panas, "
        "serta memakai validasi bergulir. Alternatif model varians bersyarat (misalnya GARCH pada "
        "residual) dapat dipertimbangkan untuk mengatasi heteroskedastisitas. Untuk keperluan "
        "kebijakan, peramalan sebaiknya disajikan sebagai rentang skenario, bukan angka tunggal.",
    )


def daftar_pustaka(document: Document) -> None:
    heading(document, "DAFTAR PUSTAKA", 1)
    references = [
        "Box, G. E. P., Jenkins, G. M., Reinsel, G. C., & Ljung, G. M. (2015). Time series analysis: Forecasting and control (5th ed.). Wiley.",
        "Cohen, A. J., Brauer, M., Burnett, R., Anderson, H. R., Frostad, J., Estep, K., ... Forouzanfar, M. H. (2017). Estimates and 25-year trends of the global burden of disease attributable to ambient air pollution: An analysis of data from the Global Burden of Diseases Study 2015. The Lancet, 389(10082), 1907–1918. https://doi.org/10.1016/S0140-6736(17)30505-6",
        "Copernicus Atmosphere Monitoring Service. (2024). CAMS global reanalysis (EAC4). ECMWF. https://ads.atmosphere.copernicus.eu/datasets/cams-global-reanalysis-eac4",
        "Hyndman, R. J., & Athanasopoulos, G. (2021). Forecasting: Principles and practice (3rd ed.). OTexts. https://otexts.com/fpp3/",
        "Open-Meteo. (2026). Air quality API documentation. https://open-meteo.com/en/docs/air-quality-api",
        "Shumway, R. H., & Stoffer, D. S. (2017). Time series analysis and its applications: With R examples (4th ed.). Springer.",
        "Wei, W. W. S. (2006). Time series analysis: Univariate and multivariate methods (2nd ed.). Pearson.",
    ]
    for reference in references:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.first_line_indent = Cm(-1.25)
        run = paragraph.add_run(reference)
        run.font.name = TNR
        run.font.size = Pt(12)


def lampiran(document: Document) -> None:
    heading(document, "LAMPIRAN A. PEMBAGIAN PERAN ANGGOTA", 1)
    add_table(
        document,
        "Pembagian peran anggota kelompok",
        ["Peran", "Anggota", "Tugas utama"],
        [
            ["Data dan eksplorasi", "[Nama Anggota 1]", "Unduh data, kontrol kualitas, EDA, dekomposisi"],
            ["Exponential Smoothing", "[Nama Anggota 2]", "MA, SES, Holt, Holt-Winters"],
            ["Keluarga ARIMA", "[Nama Anggota 3]", "AR, ARMA, ARIMA, SARIMA, Fourier"],
            ["Evaluasi dan diagnostik", "[Nama Anggota 4]", "Uji residual, perbandingan, DM"],
            ["Penulisan dan reproducibility", "[Nama Anggota 5]", "Laporan, dokumentasi, reproduksi"],
        ],
        "Penulis (2026).",
        widths=[4.0, 4.0, 7.0],
        font_size=9,
    )

    document.add_page_break()
    heading(document, "LAMPIRAN B. KODE LENGKAP", 1)
    body(
        document,
        "Kode berikut adalah versi yang dipakai untuk menghasilkan seluruh output. Pipeline utama "
        "dijalankan berurutan: fetch_openmeteo.py, preprocess.py, check_homogeneity.py, lalu "
        "analysis.py. build_report.py menghasilkan dokumen ini. Notebook walkthrough tersedia di "
        "notebooks/01_analisis_pm25_surakarta.ipynb.",
        first_line=False,
    )
    for filename, title in (
        ("fetch_openmeteo.py", "B.1 fetch_openmeteo.py"),
        ("preprocess.py", "B.2 preprocess.py"),
        ("check_homogeneity.py", "B.3 check_homogeneity.py"),
        ("analysis.py", "B.4 analysis.py"),
    ):
        add_code_listing(document, ROOT / "src" / filename, title)

    document.add_page_break()
    heading(document, "LAMPIRAN C. KELUARAN TAMBAHAN", 1)
    body(
        document,
        "Seluruh keluaran numerik tersedia di folder outputs/. Workbook Excel memuat sheet "
        "Ringkasan, Data_Mingguan, Statistik_Deskriptif, Uji_Stasioneritas, Perbandingan_Model, "
        "Grid_AR, Grid_ARMA, Grid_ARIMA, Grid_SARIMA, Grid_Fourier, SARIMA_D1, Parameter_ES, "
        "Akar_Model, Koefisien_Final, Diagnostik_Final, Peramalan, Uji_DM, Bulanan_Konteks, "
        "Rasio_PM, dan Catatan.",
        first_line=False,
    )
    add_table(
        document,
        "Daftar berkas keluaran utama",
        ["Berkas", "Isi"],
        [
            ["outputs/results.json", "Ringkasan angka analisis dalam format JSON"],
            ["outputs/Hasil_Analisis_PM25_Surakarta.xlsx", "Workbook Excel seluruh tabel"],
            ["outputs/figures/fig00–fig09", "Gambar alur, data, dekomposisi, ACF/PACF, model, diagnostik, peramalan"],
            ["outputs/tables/*.csv", "Tabel statistik, grid kandidat, koefisien, diagnostik, peramalan"],
            ["notebooks/01_analisis_pm25_surakarta.ipynb", "Walkthrough reproducible dengan output inline"],
        ],
        "Penulis (2026).",
        widths=[6.5, 8.5],
        font_size=8.5,
    )

    document.add_page_break()
    heading(document, "LAMPIRAN D. PERNYATAAN PENGGUNAAN ALAT BANTU AI", 1)
    body(
        document,
        "Alat bantu AI digunakan untuk membantu penulisan kode, merapikan struktur laporan, dan "
        "memeriksa konsistensi angka. Seluruh keputusan metodologis, pemilihan model, interpretasi "
        "hasil, dan verifikasi angka dilakukan oleh penulis. Penulis bertanggung jawab penuh atas "
        "kebenaran isi laporan dan memastikan tidak ada data yang direkayasa.",
        first_line=False,
    )


def main() -> None:
    results = load_json()
    ratios = pd.read_csv(ROOT / "data" / "processed" / "pm25_pm10_monthly.csv")
    document = setup_document()
    cover(document)
    ringkasan(document, results)
    bab1(document)
    bab2(document, results, ratios)
    bab3(document)
    bab4(document, results)
    bab5(document, results)
    daftar_pustaka(document)
    lampiran(document)
    document.save(DOCX_PATH)
    print(f"Laporan DOCX: {DOCX_PATH}")
    print(f"Jumlah tabel utama: {TABLES} | Jumlah gambar: {FIGURES}")


if __name__ == "__main__":
    main()
